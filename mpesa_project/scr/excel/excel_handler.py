import pandas as pd
from typing import List
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from scr.models.transaction import Transaction

class ExcelHandler:
    """Handle Excel file operations"""
    
    def __init__(self, excel_path: str):
        self.excel_path = Path(excel_path)
        self.ensure_file_exists()
    
    def ensure_file_exists(self):
        """Create Excel file if it doesn't exist"""
        if not self.excel_path.exists():
            df = pd.DataFrame(columns=[
                'Transaction Code', 'Date', 'Type', 'Amount', 
                'Fee', 'Sender', 'Recipient', 'Balance'
            ])
            df.to_excel(self.excel_path, index=False)
            self.format_excel()
    
    def format_excel(self):
        """Apply formatting to Excel file"""
        try:
            wb = load_workbook(self.excel_path)
            ws = wb.active
            
            # Header formatting
            header_fill = PatternFill(start_color="366092", 
                                     end_color="366092", 
                                     fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(self.excel_path)
        except Exception as e:
            print(f"Error formatting Excel: {e}")
    
    def read_existing_data(self) -> pd.DataFrame:
        """Read existing Excel data"""
        try:
            return pd.read_excel(self.excel_path)
        except Exception as e:
            print(f"Error reading Excel: {e}")
            return pd.DataFrame()
    
    def append_transactions(self, transactions: List[Transaction]) -> bool:
        """Append new transactions to Excel"""
        try:
            # Read existing data
            df_existing = self.read_existing_data()
            
            # Convert transactions to DataFrame
            new_data = []
            for trans in transactions:
                new_data.append({
                    'Transaction Code': trans.transaction_code,
                    'Date': trans.date.strftime('%Y-%m-%d %H:%M:%S'),
                    'Type': trans.transaction_type.capitalize(),
                    'Amount': trans.amount,
                    'Fee': trans.fee,
                    'Sender': trans.sender or '',
                    'Recipient': trans.recipient or '',
                    'Balance': trans.balance or ''
                })
            
            df_new = pd.DataFrame(new_data)
            
            # Remove duplicates based on transaction code
            if not df_existing.empty:
                df_combined = pd.concat([df_existing, df_new], ignore_index=True)
                df_combined = df_combined.drop_duplicates(
                    subset=['Transaction Code'], 
                    keep='first'
                )
            else:
                df_combined = df_new
            
            # Sort by date (newest first)
            df_combined['Date'] = pd.to_datetime(df_combined['Date'])
            df_combined = df_combined.sort_values('Date', ascending=False)
            
            # Save to Excel
            df_combined.to_excel(self.excel_path, index=False)
            self.format_excel()
            
            return True
            
        except Exception as e:
            print(f"Error appending to Excel: {e}")
            return False
    
    def export_from_database(self, db_manager) -> bool:
        """Export all data from database to Excel"""
        try:
            transactions = db_manager.get_all_transactions()
            
            data = []
            for trans in transactions:
                data.append({
                    'Transaction Code': trans['transaction_code'],
                    'Date': trans['date'],
                    'Type': trans['transaction_type'].capitalize(),
                    'Amount': trans['amount'],
                    'Fee': trans['fee'],
                    'Sender': trans['sender'] or '',
                    'Recipient': trans['recipient'] or '',
                    'Balance': trans['balance'] or ''
                })
            
            df = pd.DataFrame(data)
            df['Date'] = pd.to_datetime(df['Date'])
            df = df.sort_values('Date', ascending=False)
            
            df.to_excel(self.excel_path, index=False)
            self.format_excel()
            
            return True
            
        except Exception as e:
            print(f"Error exporting to Excel: {e}")
            return False